import pytest
import uuid
import json
import io
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.main import app
from app.core.database import Base, get_db
from app.core.security import get_password_hash
from app.models.organization import Organization
from app.models.user import User
from app.models.patient import Patient
from app.models.case import Case
from app.models.document import Document
from app.models.extracted_item import ExtractedItem
from app.models.audit_log import AuditLog
from app.services.ai import ai_service
from app.ocr import extract_text_from_pdf, extract_text_from_image, extract_text_from_docx, ocr_provider

from sqlalchemy.pool import StaticPool

# Let's use SQLite in-memory for fast and isolated tests
SQLALCHEMY_DATABASE_URL = "sqlite:///:memory:"

engine = create_engine(
    SQLALCHEMY_DATABASE_URL, 
    connect_args={"check_same_thread": False},
    poolclass=StaticPool
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

def override_get_db():
    try:
        db = TestingSessionLocal()
        yield db
    finally:
        db.close()

app.dependency_overrides[get_db] = override_get_db

client = TestClient(app)

@pytest.fixture(scope="module")
def setup_db():
    Base.metadata.create_all(bind=engine)
    db = TestingSessionLocal()
    
    # Org A
    org_a = Organization(name="Test Org A")
    db.add(org_a)
    db.commit()
    
    # Org B (for cross-org testing)
    org_b = Organization(name="Test Org B")
    db.add(org_b)
    db.commit()

    # Admin User A
    admin_a = User(
        org_id=org_a.id, name="Admin A", email="admin_a@test.com",
        password_hash=get_password_hash("test1234"), role="admin",
        is_active=True, email_verified=True
    )
    db.add(admin_a)
    
    # Doctor User A
    doc_a = User(
        org_id=org_a.id, name="Doc A", email="doc_a@test.com",
        password_hash=get_password_hash("test1234"), role="doctor",
        is_active=True, email_verified=True
    )
    db.add(doc_a)

    # Reviewer User A
    rev_a = User(
        org_id=org_a.id, name="Rev A", email="rev_a@test.com",
        password_hash=get_password_hash("test1234"), role="insurance_reviewer",
        is_active=True, email_verified=True
    )
    db.add(rev_a)

    staff_a = User(
        org_id=org_a.id, name="Staff A", email="staff_a@test.com",
        password_hash=get_password_hash("test1234"), role="staff",
        is_active=True, email_verified=True
    )
    db.add(staff_a)

    # Admin User B
    admin_b = User(
        org_id=org_b.id, name="Admin B", email="admin_b@test.com",
        password_hash=get_password_hash("test1234"), role="admin",
        is_active=True, email_verified=True
    )
    db.add(admin_b)
    
    db.commit()
    yield
    Base.metadata.drop_all(bind=engine)


def get_token(email: str):
    response = client.post(
        "/api/v1/auth/login",
        data={"username": email, "password": "test1234"}
    )
    return response.json()["access_token"]


def test_auth_and_org_signup(setup_db):
    response = client.post(
        "/api/v1/auth/register",
        json={
            "org_name": "New Signup Org",
            "admin_name": "New Admin",
            "admin_email": "new_admin@test.com",
            "admin_password": "newpassword123"
        }
    )
    assert response.status_code == 200
    
    login_response = client.post(
        "/api/v1/auth/login",
        data={"username": "new_admin@test.com", "password": "newpassword123"}
    )
    assert login_response.status_code == 200
    token = login_response.json()["access_token"]
    
    me_response = client.get(
        "/api/v1/auth/me",
        headers={"Authorization": f"Bearer {token}"}
    )
    assert me_response.status_code == 200
    assert me_response.json()["email"] == "new_admin@test.com"
    assert me_response.json()["role"] == "admin"


def test_rbac_matrix_and_cross_org(setup_db):
    token_admin_a = get_token("admin_a@test.com")
    token_doc_a = get_token("doc_a@test.com")
    token_rev_a = get_token("rev_a@test.com")
    token_admin_b = get_token("admin_b@test.com")

    # 1. Admin A can create a patient
    res = client.post(
        "/api/v1/patients/",
        json={"first_name": "John", "last_name": "Doe"},
        headers={"Authorization": f"Bearer {token_admin_a}"}
    )
    assert res.status_code == 200
    patient_id = res.json()["id"]

    # 2. Doctor A can create a case
    res = client.post(
        "/api/v1/cases/",
        json={"title": "Test Case", "patient_id": patient_id},
        headers={"Authorization": f"Bearer {token_doc_a}"}
    )
    assert res.status_code == 200
    case_id = res.json()["id"]

    # 3. Reviewer A can view cases but CANNOT create
    res = client.get(
        "/api/v1/cases/",
        headers={"Authorization": f"Bearer {token_rev_a}"}
    )
    assert res.status_code == 200
    
    res = client.post(
        "/api/v1/cases/",
        json={"title": "Hacked Case", "patient_id": patient_id},
        headers={"Authorization": f"Bearer {token_rev_a}"}
    )
    assert res.status_code == 403 # RBAC works

    # 4. Admin B CANNOT view Admin A's cases (cross org isolation)
    res = client.get(
        f"/api/v1/cases/{case_id}",
        headers={"Authorization": f"Bearer {token_admin_b}"}
    )
    assert res.status_code == 404

    # 5. Audit log isolation (Admin A only sees Org A's logs)
    res = client.get(
        "/api/v1/audit-logs/",
        headers={"Authorization": f"Bearer {token_admin_a}"}
    )
    assert res.status_code == 200
    assert len(res.json()) > 0
    
    res = client.get(
        "/api/v1/audit-logs/",
        headers={"Authorization": f"Bearer {token_admin_b}"}
    )
    assert res.status_code == 200
    assert len(res.json()) == 0


def test_document_upload_validation(setup_db):
    token = get_token("doc_a@test.com")
    
    # We need a case
    res = client.get("/api/v1/cases/", headers={"Authorization": f"Bearer {token}"})
    case_id = res.json()[0]["id"]
    
    # Invalid extension
    invalid_file = io.BytesIO(b"fake exe content")
    res = client.post(
        f"/api/v1/cases/{case_id}/documents",
        files={"file": ("virus.exe", invalid_file, "application/x-msdownload")},
        headers={"Authorization": f"Bearer {token}"}
    )
    assert res.status_code == 400
    assert "File type not supported" in res.json()["detail"]


def test_successful_document_upload_advances_case(setup_db, monkeypatch):
    import importlib
    documents_api = importlib.import_module("app.api.documents")
    monkeypatch.setattr(documents_api.storage_service, "save", lambda content, filename: f"tests/{filename}")
    monkeypatch.setattr(documents_api, "enqueue_job", lambda **kwargs: None)

    token = get_token("staff_a@test.com")
    headers = {"Authorization": f"Bearer {token}"}
    patient = client.post(
        "/api/v1/patients/",
        json={"first_name": "Upload", "last_name": "Patient", "dob": "1991-01-01"},
        headers=headers,
    ).json()
    case = client.post(
        "/api/v1/cases/",
        json={"patient_id": patient["id"], "title": "Upload Test"},
        headers=headers,
    ).json()
    response = client.post(
        f"/api/v1/cases/{case['id']}/documents",
        files={"file": ("record.pdf", io.BytesIO(b"%PDF-1.4 synthetic"), "application/pdf")},
        headers=headers,
    )
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "UPLOADED"
    refreshed = client.get(f"/api/v1/cases/{case['id']}", headers=headers)
    assert refreshed.json()["status"] == "DOCUMENTS_UPLOADED"


def test_new_case_workflow_new_and_existing_patient(setup_db):
    token = get_token("staff_a@test.com")
    headers = {"Authorization": f"Bearer {token}"}

    users = client.get("/api/v1/users/?role=doctor&is_active=true", headers=headers)
    assert users.status_code == 200
    assert users.json() and all(user["role"] == "doctor" for user in users.json())
    doctor_id = users.json()[0]["id"]

    no_insurance = client.post(
        "/api/v1/cases/",
        json={
            "patient": {
                "first_name": "John", "last_name": "Doe", "dob": "1990-01-01",
                "gender": "male", "identifiers": {"mrn": "MRN-JOHN-1"}
            },
            "title": "General Consultation",
            "diagnoses": {"primary_condition": "Other"},
            "symptoms": {"reason_for_visit": "Headache"},
            "insurance_available": False,
            "assigned_to": doctor_id,
            "additional_notes": "Patient referred for evaluation."
        },
        headers=headers,
    )
    assert no_insurance.status_code == 200, no_insurance.text
    first_case = no_insurance.json()
    assert first_case["status"] == "NEW"
    assert first_case["insurance_available"] is False
    assert first_case["insurance_provider"] is None
    patient_id = first_case["patient_id"]

    before = TestingSessionLocal().query(Patient).count()
    existing = client.post(
        "/api/v1/cases/",
        json={"patient_id": patient_id, "title": "Follow-up", "insurance_available": False},
        headers=headers,
    )
    assert existing.status_code == 200, existing.text
    db = TestingSessionLocal()
    assert db.query(Patient).count() == before
    assert db.query(AuditLog).filter(AuditLog.action == "CASE_CREATED", AuditLog.entity_id == str(first_case["id"])).one()
    db.close()


def test_new_case_insurance_and_duplicate_protection(setup_db):
    token = get_token("staff_a@test.com")
    headers = {"Authorization": f"Bearer {token}"}
    reviewers = client.get("/api/v1/users/?role=insurance_reviewer&is_active=true", headers=headers)
    reviewer_id = reviewers.json()[0]["id"]

    payload = {
        "patient": {"first_name": "Jane", "last_name": "Smith", "dob": "1985-06-10"},
        "title": "Hypertension Follow-up",
        "insurance_available": True,
        "insurance_provider": "Example Provider",
        "insurance_number": "POL123456",
        "assigned_insurance_reviewer": reviewer_id,
    }
    created = client.post("/api/v1/cases/", json=payload, headers=headers)
    assert created.status_code == 200, created.text
    assert created.json()["insurance_provider"] == "Example Provider"
    assert created.json()["assigned_insurance_reviewer"] == reviewer_id

    db = TestingSessionLocal()
    patient_count = db.query(Patient).filter(Patient.first_name == "Jane", Patient.last_name == "Smith").count()
    case_count = db.query(Case).filter(Case.title == "Hypertension Follow-up").count()
    db.close()

    duplicate = client.post("/api/v1/cases/", json=payload, headers=headers)
    assert duplicate.status_code == 409
    db = TestingSessionLocal()
    assert db.query(Patient).filter(Patient.first_name == "Jane", Patient.last_name == "Smith").count() == patient_count
    assert db.query(Case).filter(Case.title == "Hypertension Follow-up").count() == case_count
    db.close()


def test_staff_cannot_review_and_required_reviews_complete_case(setup_db):
    staff_token = get_token("staff_a@test.com")
    doctor_token = get_token("doc_a@test.com")
    reviewer_token = get_token("rev_a@test.com")
    db = TestingSessionLocal()
    doctor = db.query(User).filter(User.email == "doc_a@test.com").one()
    reviewer = db.query(User).filter(User.email == "rev_a@test.com").one()
    db.close()

    created = client.post(
        "/api/v1/cases/",
        json={
            "patient": {"first_name": "Review", "last_name": "Patient", "dob": "1975-02-03"},
            "title": "Review workflow",
            "insurance_available": True,
            "insurance_provider": "Example Provider",
            "insurance_number": "REVIEW-123",
            "assigned_to": doctor.id,
            "assigned_insurance_reviewer": reviewer.id,
        },
        headers={"Authorization": f"Bearer {staff_token}"},
    )
    case_id = created.json()["id"]
    db = TestingSessionLocal()
    case = db.query(Case).filter(Case.id == case_id).one()
    case.status = "UNDER_REVIEW"
    db.commit()
    db.close()

    forbidden = client.post(
        f"/api/v1/cases/{case_id}/review",
        json={"decision": "APPROVED"},
        headers={"Authorization": f"Bearer {staff_token}"},
    )
    assert forbidden.status_code == 403

    doctor_approved = client.post(
        f"/api/v1/cases/{case_id}/review",
        json={"decision": "APPROVED"},
        headers={"Authorization": f"Bearer {doctor_token}"},
    )
    assert doctor_approved.status_code == 200
    assert doctor_approved.json()["status"] == "UNDER_REVIEW"

    reviewer_approved = client.post(
        f"/api/v1/cases/{case_id}/review",
        json={"decision": "APPROVED"},
        headers={"Authorization": f"Bearer {reviewer_token}"},
    )
    assert reviewer_approved.status_code == 200
    assert reviewer_approved.json()["status"] == "COMPLETED"

def test_llm_mock(setup_db):
    # Test our mock AI service fallback
    old_provider = ai_service.provider
    ai_service.provider = None
    try:
        result = ai_service.generate_analysis(1, "Patient has a headache.")
        assert result.confidence == "low"
        assert "Mock summary" in result.patient_summary
    finally:
        ai_service.provider = old_provider


def test_pdf_and_docx_text_extraction():
    import fitz
    from docx import Document as DocxDocument

    pdf = fitz.open()
    page = pdf.new_page()
    expected_pdf_text = "Synthetic clinical document with enough plain text for direct PDF extraction."
    page.insert_text((72, 72), expected_pdf_text)
    pages, full_text = extract_text_from_pdf(pdf.tobytes())
    pdf.close()
    assert len(pages) == 1
    assert "Synthetic clinical document" in full_text

    docx = DocxDocument()
    docx.add_paragraph("Synthetic DOCX medical document")
    buffer = io.BytesIO()
    docx.save(buffer)
    pages, full_text = extract_text_from_docx(buffer.getvalue())
    assert pages == ["Synthetic DOCX medical document"]
    assert full_text == "Synthetic DOCX medical document"


def test_scanned_image_ocr_when_runtime_available():
    if type(ocr_provider).__name__ == "UnavailableOCRProvider":
        pytest.skip("Tesseract runtime is not installed")
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1200, 260), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 72)
    draw.text((40, 70), "MEDICAL OCR TEST 12345", font=font, fill="black")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    _, text = extract_text_from_image(buffer.getvalue())
    assert "MEDICAL OCR TEST" in text.upper()
    assert "12345" in text


def test_review_confirm_edit_reject(setup_db):
    db = TestingSessionLocal()
    org_a = db.query(Organization).filter(Organization.name == "Test Org A").first()
    
    patient = Patient(org_id=org_a.id, first_name="Test", last_name="Patient", created_by=1)
    db.add(patient)
    db.commit()
    
    case = Case(org_id=org_a.id, patient_id=patient.id, title="Test Case", status="NEW", created_by=1)
    db.add(case)
    db.commit()
    
    doc = Document(org_id=org_a.id, case_id=case.id, filename="x", mime_type="x", size=1, doc_type="x", storage_key="x", uploaded_by=1)
    db.add(doc)
    db.commit()
    db.refresh(doc)
    
    item = ExtractedItem(org_id=org_a.id, case_id=case.id, document_id=doc.id, category="diagnosis", value="Fever", review_status="PENDING", ai_generated=True)
    db.add(item)
    db.commit()
    db.refresh(item)
    item_id = item.id
    db.close()

    token_doc = get_token("doc_a@test.com")
    
    # Edit
    res = client.patch(
        f"/api/v1/extracted-items/{item_id}",
        json={"action": "edit", "value": "High Fever"},
        headers={"Authorization": f"Bearer {token_doc}"}
    )
    assert res.status_code == 200
    assert res.json()["review_status"] == "EDITED"
    assert res.json()["value"] == "High Fever"
    assert res.json()["original_value"] == "Fever"

    # Confirm
    res = client.patch(
        f"/api/v1/extracted-items/{item_id}",
        json={"action": "confirm"},
        headers={"Authorization": f"Bearer {token_doc}"}
    )
    assert res.status_code == 200
    assert res.json()["review_status"] == "CONFIRMED"

    # Reject
    res = client.patch(
        f"/api/v1/extracted-items/{item_id}",
        json={"action": "reject"},
        headers={"Authorization": f"Bearer {token_doc}"}
    )
    assert res.status_code == 200
    assert res.json()["review_status"] == "REJECTED"

